#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""meeting-minutes v2 · Phase 5 出 PPT（按需）

读取 `纪要/演示大纲.json`（由模型根据已通过质检的纪要撰写），生成
`纪要/会议汇报.pptx`。页面顺序与配色见 references/output-formats.md。

为什么用大纲 JSON 而不是直接解析纪要：PPT 的取舍（哪 5 条最该上首页、
哪两个议题合并成一页）是判断题，交给模型；排版是确定性的活，交给脚本。

大纲格式：
{
  "title": "会议名 纪要汇报",
  "subtitle": "2026-09-19 ｜ 参会人",
  "slides": [
    {"type": "bullets",     "title": "会议背景与目标", "items": ["...", "..."]},
    {"type": "conclusions", "title": "核心结论速览",
     "items": [{"text": "...", "status": "已确认"}]},
    {"type": "table", "title": "行动项",
     "header": ["事项", "责任人", "截止日"],
     "rows": [["...", "...", "..."]]},
    {"type": "note", "title": "待确认与风险", "items": ["..."]}
  ]
}

用法：
    python build_ppt.py "<工作目录>"
    python build_ppt.py "<工作目录>" --spec "<工作目录>/纪要/演示大纲.json"
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

MAIN = "5B7C99"
LIGHT = "EDF2F7"
GRAY = "8A8F98"
RED = "B4646B"
GREEN = "4E7C6A"
GOLD = "A08A63"
FONT = "微软雅黑"

STATUS_COLOR = {"已确认": MAIN, "已决策": MAIN, "待确认": GRAY,
                "已否决": RED, "已延期": GOLD, "风险": RED}


def rgb(hexstr):
    from pptx.dml.color import RGBColor
    h = hexstr.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def set_font(run, size=None, bold=None, color=None, name=FONT):
    from pptx.util import Pt
    from pptx.oxml.ns import qn
    run.font.name = name
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", name)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def rect(slide, x, y, w, h, fill):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def textbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    return tf


def add_header(slide, prs, title, page=None):
    from pptx.util import Inches, Pt
    rect(slide, Inches(0), Inches(0), prs.slide_width, Inches(1.05), MAIN)
    tf = textbox(slide, Inches(0.55), Inches(0.18), prs.slide_width - Inches(1.6), Inches(0.7))
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_font(r, 26, True, "FFFFFF")
    if page is not None:
        tf2 = textbox(slide, prs.slide_width - Inches(1.25), Inches(0.3), Inches(0.9), Inches(0.5))
        p2 = tf2.paragraphs[0]
        r2 = p2.add_run()
        r2.text = str(page)
        set_font(r2, 14, False, "FFFFFF")


def add_bullets(slide, prs, items, top=1.45, size=18):
    from pptx.util import Inches, Pt
    tf = textbox(slide, Inches(0.75), Inches(top), prs.slide_width - Inches(1.5),
                 prs.slide_height - Inches(top + 0.5))
    for k, it in enumerate(items):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        r = p.add_run()
        r.text = "·  " + str(it)
        set_font(r, size, False, "2B2F36")
    return tf


def add_conclusions(slide, prs, items, top=1.4):
    from pptx.util import Inches, Pt
    y = Inches(top)
    for it in items[:5]:
        if isinstance(it, str):
            text, status = it, ""
        else:
            text, status = it.get("text", ""), it.get("status", "")
        h = Inches(0.92)
        color = STATUS_COLOR.get(status, MAIN)
        rect(slide, Inches(0.6), y, Inches(0.09), h - Inches(0.14), color)
        tf = textbox(slide, Inches(0.85), y, prs.slide_width - Inches(2.6), h)
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = text
        set_font(r, 17, False, "2B2F36")
        if status:
            tfs = textbox(slide, prs.slide_width - Inches(1.75), y + Inches(0.18),
                          Inches(1.2), Inches(0.5))
            ps = tfs.paragraphs[0]
            rs = ps.add_run()
            rs.text = status
            set_font(rs, 12, True, color)
        y += h


def add_table(slide, prs, header, rows, top=1.45):
    from pptx.util import Inches, Pt
    cols = len(header)
    shape = slide.shapes.add_table(len(rows) + 1, cols, Inches(0.6), Inches(top),
                                   prs.slide_width - Inches(1.2),
                                   Inches(min(4.6, 0.5 + 0.46 * len(rows))))
    table = shape.table
    for k, htxt in enumerate(header):
        cell = table.cell(0, k)
        cell.text = ""
        p = cell.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = str(htxt)
        set_font(r, 13, True, "FFFFFF")
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(MAIN)
    for ri, row in enumerate(rows, start=1):
        for k in range(cols):
            cell = table.cell(ri, k)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            r = p.add_run()
            val = str(row[k]) if k < len(row) else ""
            r.text = val
            color = GRAY if val.strip() in ("待确认", "—", "") else "2B2F36"
            set_font(r, 12, False, color)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb("FFFFFF" if ri % 2 else "F7FAFC")


def add_cover(prs, spec):
    from pptx.util import Inches, Pt
    slide = blank(prs)
    rect(slide, Inches(0), Inches(0), prs.slide_width, prs.slide_height, MAIN)
    rect(slide, Inches(0), Inches(4.15), prs.slide_width, Inches(0.06), "FFFFFF")
    tf = textbox(slide, Inches(1.1), Inches(2.3), prs.slide_width - Inches(2.2), Inches(1.6))
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = spec.get("title", "会议纪要")
    set_font(r, 40, True, "FFFFFF")
    tf2 = textbox(slide, Inches(1.1), Inches(4.5), prs.slide_width - Inches(2.2), Inches(1.6))
    for k, line in enumerate(spec.get("subtitle_lines") or [spec.get("subtitle", "")]):
        p = tf2.paragraphs[0] if k == 0 else tf2.add_paragraph()
        r = p.add_run()
        r.text = str(line)
        set_font(r, 16, False, "E8EFF5")
    return slide


def build(spec: dict, out_path: Path) -> int:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_cover(prs, spec)

    page = 2
    for s in spec.get("slides", []):
        stype = s.get("type", "bullets")
        slide = blank(prs)
        add_header(slide, prs, s.get("title", ""), page)
        if stype == "bullets":
            add_bullets(slide, prs, s.get("items", []))
        elif stype == "conclusions":
            add_conclusions(slide, prs, s.get("items", []))
        elif stype == "table":
            add_table(slide, prs, s.get("header", []), s.get("rows", []))
        elif stype == "note":
            add_bullets(slide, prs, s.get("items", []), size=16)
        else:
            add_bullets(slide, prs, s.get("items", []))
        page += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return len(prs.slides._sldIdLst)     # noqa: SLF001


def main() -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

    ap = argparse.ArgumentParser(description="meeting-minutes v2 出 PPT")
    ap.add_argument("workdir")
    ap.add_argument("--spec", default="", help="大纲 JSON 路径")
    ap.add_argument("--out", default="", help="输出 pptx 路径")
    args = ap.parse_args()

    wd = Path(args.workdir).expanduser().resolve()
    spec_path = Path(args.spec).expanduser() if args.spec else wd / "纪要" / "演示大纲.json"
    if not spec_path.exists():
        print(f"[错误] 找不到大纲：{spec_path}\n"
              f"       先写一份《演示大纲.json》，再跑本脚本。", file=sys.stderr)
        return 2
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    out = Path(args.out).expanduser() if args.out else wd / "纪要" / "会议汇报.pptx"
    n = build(spec, out)
    print(f"[生成] {out}（{n} 页）")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception:
        import traceback
        traceback.print_exc()
        sys.exit(1)
