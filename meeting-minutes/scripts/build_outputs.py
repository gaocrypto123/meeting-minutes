#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""meeting-minutes v2 · Phase 5 出稿：Markdown → HTML / Word（卡片式排版）

从纪要母版派生可直接查看的成品：
    <工作目录>/原文/转写稿.html     带锚点，供跳转定位
    <工作目录>/纪要/*.html          带目录、卡片式条目；[时:分:秒] 可点回原文
    <工作目录>/纪要/会议纪要.docx    Word 归档版，同样分层渲染

排版规则：
    * H1 之后、第一个 H2 之前的内容 → 「速览」高亮块
    * 每条 `- **类型 ID**：结论` + 缩进字段 → 一张卡片：类型色标 + 结论 + 证据条 + 引用
    * 引用类字段单独成引文块，其余字段压缩成一行证据条（不再一行一个字段）

用法：
    python build_outputs.py "<工作目录>"
    python build_outputs.py "<工作目录>" --no-docx
"""
from __future__ import annotations

import argparse
import html
import re
import sys
from pathlib import Path

TS_RE = re.compile(r"\[(\d{2}):(\d{2}):(\d{2})\]")
TURN_RE = re.compile(r"^\d+\.\s*\[(\d{2}):(\d{2}):(\d{2})\]\s*(\S+?)[：:]\s*(.*)$")
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
CODE_RE = re.compile(r"`([^`]+)`")
ENTRY_RE = re.compile(r"^-\s*\*\*(?P<kind>[\u4e00-\u9fff]{2,6})\s*(?P<id>[A-Za-z]{0,3}\d{0,4})\*\*"
                      r"\s*[：:]\s*(?P<title>.*)$")
FIELD_RE = re.compile(r"^\s+-\s*(?P<key>[\u4e00-\u9fff]{2,6})\s*[：:]\s*(?P<val>.*)$")

KIND_COLOR = {
    "行动项": "#4E7C6A", "风险": "#B4646B", "问题": "#B4646B", "待定事项": "#B4646B",
}
STATUS_COLOR = {
    "已确认": "#5B7C99", "已决策": "#5B7C99", "待确认": "#8A8F98",
    "已否决": "#B4646B", "已延期": "#A08A63",
}
QUOTE_FIELDS = {"引用"}

CSS = """
:root { --main:#5B7C99; --light:#EDF2F7; --gray:#8A8F98; --red:#B4646B; --green:#4E7C6A; }
* { box-sizing: border-box; }
body { font-family:"Microsoft YaHei","PingFang SC","Segoe UI",sans-serif; line-height:1.75;
  color:#2B2F36; max-width:920px; margin:0 auto; padding:44px 28px 90px; background:#fff; }
h1 { font-size:27px; border-bottom:3px solid var(--main); padding-bottom:10px; margin-bottom:8px; }
h2 { font-size:20px; margin:38px 0 14px; padding:6px 0 6px 13px; border-left:5px solid var(--main);
  background:linear-gradient(90deg,#F5F8FB,transparent); }
h3 { font-size:16px; color:var(--main); margin:26px 0 10px; }
h4 { font-size:15px; color:#4A5461; margin:20px 0 8px; }
blockquote { margin:14px 0; padding:12px 16px; background:var(--light);
  border-left:4px solid var(--main); color:#47505C; font-size:14px; border-radius:0 6px 6px 0; }
table { border-collapse:collapse; width:100%; margin:16px 0; font-size:14px; }
th { background:var(--main); color:#fff; text-align:left; font-weight:600; }
th,td { border:1px solid #D8E0E8; padding:8px 11px; vertical-align:top; }
tr:nth-child(even) td { background:#F7FAFC; }
ul,ol { padding-left:24px; } li { margin:5px 0; }
code { background:#F2F5F8; padding:1px 5px; border-radius:3px; font-size:.94em; }
a.ts { color:var(--main); text-decoration:none; font-variant-numeric:tabular-nums; }
a.ts:hover { text-decoration:underline; }
.lead { background:#F5F8FB; border:1px solid #DCE6EF; border-radius:8px;
  padding:6px 20px 14px; margin:18px 0 8px; }
.lead h3 { margin-top:14px; }
.toc { background:#FAFCFD; border:1px dashed #CBD8E3; border-radius:8px;
  padding:12px 20px; margin:20px 0 8px; font-size:14px; }
.toc b { color:var(--main); } .toc a { color:#3F5C74; text-decoration:none; }
.toc a:hover { text-decoration:underline; } .toc ol { margin:6px 0 0; padding-left:20px; }
.entry { border:1px solid #E1E9F0; border-left:4px solid var(--main); border-radius:6px;
  padding:11px 15px; margin:11px 0; background:#FDFEFE; }
.badge { display:inline-block; background:var(--main); color:#fff; font-size:12px;
  padding:1px 9px; border-radius:11px; margin-right:7px; vertical-align:1px; }
.eid { color:var(--gray); font-size:12.5px; margin-right:6px; font-variant-numeric:tabular-nums; }
.entry-title { margin:0; font-weight:600; }
.meta { margin:7px 0 0; font-size:12.5px; color:#5A6472; }
.chip { display:inline-block; background:#EEF3F8; border-radius:4px; padding:1px 8px;
  margin:0 6px 3px 0; }
.chip.ok { background:#E8F0F7; color:#3F5C74; }
.chip.wait { background:#F0F1F3; color:#6C737D; }
.chip.bad { background:#F7ECED; color:#9A4F55; }
.quote { margin:8px 0 0; padding:7px 12px; background:#F7F9FB; border-left:3px solid #C9D6E2;
  font-size:13px; color:#4A5461; border-radius:0 4px 4px 0; }
.turn { margin:6px 0; }
.turn .spk { color:var(--main); font-weight:600; }
hr { border:none; border-top:1px solid #E3E9EF; margin:30px 0; }
"""


def esc_bold_code(text: str, link_ts: str | None = None) -> str:
    out = html.escape(text, quote=False)
    out = BOLD_RE.sub(r"<strong>\1</strong>", out)
    out = CODE_RE.sub(r"<code>\1</code>", out)
    if link_ts:
        out = TS_RE.sub(
            lambda m: f'<a class="ts" href="{link_ts}#t-{m.group(1)}{m.group(2)}{m.group(3)}">'
                      f'[{m.group(1)}:{m.group(2)}:{m.group(3)}]</a>', out)
    return out


def status_class(value: str) -> str:
    v = value.strip()
    if v in ("已确认", "已决策"):
        return "chip ok"
    if v in ("已否决",):
        return "chip bad"
    return "chip wait"


def render_entry(kind: str, eid: str, title: str, fields: list,
                 link_ts: str | None) -> str:
    color = KIND_COLOR.get(kind, "#5B7C99")
    chips, quote = [], ""
    for key, val in fields:
        if key in QUOTE_FIELDS:
            quote = f'<div class="quote">{esc_bold_code(val, link_ts)}</div>'
        elif key == "状态":
            chips.append(f'<span class="{status_class(val)}">状态：{html.escape(val)}</span>')
        else:
            chips.append(f'<span class="chip">{html.escape(key)}：'
                         f'{esc_bold_code(val, link_ts)}</span>')
    head = (f'<span class="badge" style="background:{color}">{html.escape(kind)}</span>'
            + (f'<span class="eid">{html.escape(eid)}</span>' if eid else ""))
    return (f'<div class="entry" style="border-left-color:{color}">'
            f'<p class="entry-title">{head}{esc_bold_code(title, link_ts)}</p>'
            + (f'<p class="meta">{"".join(chips)}</p>' if chips else "")
            + quote + "</div>")


def md_to_html(md: str, title: str, link_ts: str | None = None,
               transcript: bool = False) -> str:
    lines = md.splitlines()
    lead, body, toc = [], [], []
    seen_h1 = seen_h2 = False
    i = 0

    def flush_lead():
        if lead:
            body.append('<div class="lead">' + "".join(lead) + "</div>")
            lead.clear()

    while i < len(lines):
        line = lines[i].rstrip()
        s = line.strip()
        if not s:
            i += 1
            continue

        m = ENTRY_RE.match(s)
        if m:
            fields = []
            j = i + 1
            while j < len(lines):
                fm = FIELD_RE.match(lines[j])
                if not fm:
                    break
                fields.append((fm.group("key"), fm.group("val").strip()))
                j += 1
            node = render_entry(m.group("kind"), (m.group("id") or "").strip(),
                                m.group("title").strip(), fields, link_ts)
            (lead if (seen_h1 and not seen_h2) else body).append(node)
            i = j
            continue

        if s.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            body.append("<pre><code>" + html.escape("\n".join(buf)) + "</code></pre>")
            continue
        if re.fullmatch(r"-{3,}", s):
            body.append("<hr>")
            i += 1
            continue
        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-|]+\|$",
                                                               lines[i + 1].strip()):
            header = [c.strip() for c in s.strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            node = ("<table><thead><tr>"
                    + "".join(f"<th>{esc_bold_code(c, link_ts)}</th>" for c in header)
                    + "</tr></thead><tbody>"
                    + "".join("<tr>" + "".join(f"<td>{esc_bold_code(c, link_ts)}</td>"
                                               for c in r) + "</tr>" for r in rows)
                    + "</tbody></table>")
            (lead if (seen_h1 and not seen_h2) else body).append(node)
            continue

        handled = False
        for lvl, prefix in ((1, "# "), (2, "## "), (3, "### "), (4, "#### ")):
            if s.startswith(prefix):
                txt = s[len(prefix):]
                if lvl == 1:
                    seen_h1 = True
                elif lvl == 2:
                    if not seen_h2:
                        flush_lead()
                    seen_h2 = True
                    toc.append(txt)
                node = f"<h{lvl}>{esc_bold_code(txt, link_ts)}</h{lvl}>"
                if lvl == 1:
                    body.append(node)
                elif lvl == 2:
                    body.append(f'<h2 id="sec-{len(toc)}">{esc_bold_code(txt, link_ts)}</h2>')
                else:
                    (lead if not seen_h2 else body).append(node)
                handled = True
                break
        if handled:
            i += 1
            continue

        if s.startswith(">"):
            node = f"<blockquote>{esc_bold_code(s.lstrip('> '), link_ts)}</blockquote>"
        elif transcript and TURN_RE.match(s):
            m = TURN_RE.match(s)
            t = f"t-{m.group(1)}{m.group(2)}{m.group(3)}"
            node = (f'<p class="turn" id="{t}">'
                    f'<span class="ts">[{m.group(1)}:{m.group(2)}:{m.group(3)}]</span> '
                    f'<span class="spk">{html.escape(m.group(4))}</span>：'
                    f'{html.escape(m.group(5))}</p>')
        elif re.match(r"^\d+\.\s", s):
            items = []
            while i < len(lines) and re.match(r"^\d+\.\s", lines[i].strip()):
                items.append(re.sub(r"^\d+\.\s*", "", lines[i].strip()))
                i += 1
            node = "<ol>" + "".join(f"<li>{esc_bold_code(x, link_ts)}</li>"
                                    for x in items) + "</ol>"
            (lead if (seen_h1 and not seen_h2) else body).append(node)
            continue
        elif s.startswith("- "):
            items = []
            while i < len(lines) and lines[i].strip().startswith("- "):
                raw = lines[i]
                items.append(((len(raw) - len(raw.lstrip())) // 2, raw.strip()[2:]))
                i += 1
            node = "<ul>" + "".join(
                (f'<li style="margin-left:{d * 22}px">' if d else "<li>")
                + esc_bold_code(t, link_ts) + "</li>" for d, t in items) + "</ul>"
            (lead if (seen_h1 and not seen_h2) else body).append(node)
            continue
        else:
            node = f"<p>{esc_bold_code(s, link_ts)}</p>"
        (lead if (seen_h1 and not seen_h2) else body).append(node)
        i += 1

    flush_lead()
    toc_html = ""
    if toc:
        toc_html = ('<div class="toc"><b>目录</b><ol>'
                    + "".join(f'<li><a href="#sec-{k + 1}">{html.escape(t)}</a></li>'
                              for k, t in enumerate(toc))
                    + "</ol></div>")
        # 插到 H1 之后
        for k, node in enumerate(body):
            if node.startswith("<h1"):
                body.insert(k + 1, toc_html)
                break
    return ('<!doctype html>\n<html lang="zh-CN"><head><meta charset="utf-8">'
            f"<title>{html.escape(title)}</title><style>{CSS}</style></head><body>"
            + "".join(body) + "</body></html>")


# ---------------------------------------------------------------- Word

def md_to_docx(md: str, out_path: Path, title: str) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    from docx.oxml import OxmlElement

    def shade(par, hexcolor: str):
        pPr = par._p.get_or_add_pPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:fill"), hexcolor)
        pPr.append(sh)

    def left_bar(par, hexcolor: str):
        pPr = par._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "6")
        left.set(qn("w:color"), hexcolor.lstrip("#"))
        bdr.append(left)
        pPr.append(bdr)

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "微软雅黑"
    st.font.size = Pt(10.5)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    def add_runs(par, text, bold_all=False, size=None, color=None):
        for k, seg in enumerate(BOLD_RE.split(text)):
            if not seg:
                continue
            run = par.add_run(seg)
            run.bold = bold_all or (k % 2 == 1)
            if size:
                run.font.size = size
            if color:
                run.font.color.rgb = color
        return par

    lines = md.splitlines()
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        if re.fullmatch(r"-{3,}", s):
            i += 1
            continue

        m = ENTRY_RE.match(s)
        if m:
            fields = []
            j = i + 1
            while j < len(lines):
                fm = FIELD_RE.match(lines[j])
                if not fm:
                    break
                fields.append((fm.group("key"), fm.group("val").strip()))
                j += 1
            kind = m.group("kind")
            color = RGBColor.from_string(KIND_COLOR.get(kind, "#5B7C99").lstrip("#"))
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            left_bar(p, KIND_COLOR.get(kind, "#5B7C99"))
            tag = f"【{kind}{(m.group('id') or '').strip()}】"
            r = p.add_run(tag)
            r.bold = True
            r.font.color.rgb = color
            add_runs(p, m.group("title").strip(), bold_all=True)
            chips = [(k, v) for k, v in fields if k not in QUOTE_FIELDS]
            if chips:
                mp = doc.add_paragraph()
                mp.paragraph_format.space_after = Pt(2)
                joined = "   ｜   ".join(
                    "{}：{}".format(k, BOLD_RE.sub(r"\1", v)) for k, v in chips)
                r = mp.add_run(joined)
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(0x6C, 0x73, 0x7D)
            for k, v in fields:
                if k in QUOTE_FIELDS:
                    qp = doc.add_paragraph()
                    qp.paragraph_format.left_indent = Pt(16)
                    qp.paragraph_format.space_after = Pt(2)
                    r = qp.add_run(BOLD_RE.sub(r"\1", v))
                    r.italic = True
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(0x4A, 0x54, 0x61)
            i = j
            continue

        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-|]+\|$",
                                                                lines[i + 1].strip()):
            header = [c.strip() for c in s.strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            for k, h in enumerate(header):
                cell = table.rows[0].cells[k]
                cell.text = ""
                r = cell.paragraphs[0].add_run(h)
                r.bold = True
                r.font.color.rgb = RGBColor(0x5B, 0x7C, 0x99)
            for row in rows:
                cells = table.add_row().cells
                for k, v in enumerate(row[:len(header)]):
                    cells[k].text = BOLD_RE.sub(r"\1", v)
            continue
        if s.startswith("#### "):
            doc.add_heading(BOLD_RE.sub(r"\1", s[5:]), level=4)
        elif s.startswith("### "):
            doc.add_heading(BOLD_RE.sub(r"\1", s[4:]), level=3)
        elif s.startswith("## "):
            doc.add_heading(BOLD_RE.sub(r"\1", s[3:]), level=2)
        elif s.startswith("# "):
            h = doc.add_heading(BOLD_RE.sub(r"\1", s[2:]), level=0)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0x5B, 0x7C, 0x99)
        elif s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(16)
            shade(p, "EDF2F7")
            r = p.add_run(s.lstrip("> "))
            r.italic = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x47, 0x50, 0x5C)
        elif s.startswith("- "):
            while i < len(lines) and lines[i].strip().startswith("- "):
                raw = lines[i]
                depth = (len(raw) - len(raw.lstrip())) // 2
                p = doc.add_paragraph(style="List Bullet" if depth == 0 else "List Bullet 2")
                add_runs(p, raw.strip()[2:])
                i += 1
            continue
        elif re.match(r"^\d+\.\s", s):
            while i < len(lines) and re.match(r"^\d+\.\s", lines[i].strip()):
                p = doc.add_paragraph(style="List Number")
                add_runs(p, re.sub(r"^\d+\.\s*", "", lines[i].strip()))
                i += 1
            continue
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs(p, s)
        i += 1

    doc.save(str(out_path))


def main() -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

    ap = argparse.ArgumentParser(description="meeting-minutes v2 出稿")
    ap.add_argument("workdir", help="会议工作目录")
    ap.add_argument("--no-docx", action="store_true", help="只出 HTML")
    args = ap.parse_args()

    wd = Path(args.workdir).expanduser().resolve()
    src_md = wd / "原文" / "转写稿.md"
    min_dir = wd / "纪要"
    if not min_dir.exists():
        print(f"[错误] 找不到纪要目录：{min_dir}", file=sys.stderr)
        return 2

    made = []
    if src_md.exists():
        out = src_md.with_suffix(".html")
        out.write_text(md_to_html(src_md.read_text(encoding="utf-8"),
                                  "转写稿 · 原文", transcript=True), encoding="utf-8")
        made.append(out)

    for md in sorted(min_dir.glob("*.md")):
        out = md.with_suffix(".html")
        out.write_text(md_to_html(md.read_text(encoding="utf-8"), md.stem,
                                  link_ts="../原文/转写稿.html"), encoding="utf-8")
        made.append(out)

    main_md = min_dir / "会议纪要.md"
    if not args.no_docx and main_md.exists():
        target = min_dir / "会议纪要.docx"
        md_to_docx(main_md.read_text(encoding="utf-8"), target, "会议纪要")
        made.append(target)

    for p in made:
        print(f"[生成] {p}")
    print(f"\n共生成 {len(made)} 个文件。HTML 双击查看；纪要里 [时:分:秒] 可点回原文。")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception:
        import traceback
        traceback.print_exc()
        sys.exit(1)
